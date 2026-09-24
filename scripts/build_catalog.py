#!/usr/bin/env python3
"""Build js/catalog.js from Excel + local seed images. qc_batches stays empty until agent feed lands."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SHEETS_ROOT = ROOT.parent / "kakobuyqcsheets"
XLSX = SHEETS_ROOT / "2026-09-10.xlsx"
PRODUCTS_DIR = SHEETS_ROOT / "img" / "products"
OUT = ROOT / "js" / "catalog.js"

CNY_TO_USD = 7.2

CAT_MAP = {
    "shoes": "shoes",
    "jordan 4": "shoes",
    "tops": "tops",
    "polo": "tops",
    "shirts": "tops",
    "shirt": "tops",
    "hoodies": "hoodies",
    "hoodie": "hoodies",
    "nike tech": "hoodies",
    "jacket": "jackets",
    "jackets": "jackets",
    "bottoms": "bottoms",
    "shorts": "bottoms",
    "pants": "bottoms",
    "bag": "bags",
    "bags": "bags",
    "accessories": "accessories",
    "underwear": "underwear",
    "watch": "watches",
    "watches": "watches",
    "electronics": "electronics",
    "jersey": "jersey",
    "set": "sets",
    "sets": "sets",
    "tracksuit": "sets",
    "suit": "sets",
    "other": "other",
}

CATEGORIES = [
    {"slug": "shoes", "label": "Shoes"},
    {"slug": "tops", "label": "Tops"},
    {"slug": "hoodies", "label": "Hoodies"},
    {"slug": "jackets", "label": "Jackets"},
    {"slug": "bottoms", "label": "Bottoms"},
    {"slug": "bags", "label": "Bags"},
    {"slug": "accessories", "label": "Accessories"},
    {"slug": "sets", "label": "Sets"},
    {"slug": "jersey", "label": "Jersey"},
    {"slug": "watches", "label": "Watches"},
    {"slug": "underwear", "label": "Underwear"},
    {"slug": "electronics", "label": "Electronics"},
    {"slug": "other", "label": "Other"},
]

# Prefer clean white-bg cutouts for home category icons
CAT_IMAGE_OVERRIDES = {
    "tops": "img/products/7725725581.webp",
    "bottoms": "img/products/7730552342.webp",
    "shoes": "img/products/7728700946.webp",  # Jordan 3, white bg
    "accessories": "img/products/7730242208.webp",  # belt, white bg
    "jackets": "img/products/7797943764.webp",  # Nike jacket, white bg
    "bags": "img/products/7730179090.webp",  # LV tote, white bg
    "watches": "img/products/7728742916.webp",  # Apple Watch collage, white bg
    "electronics": "img/products/7730271602.webp",  # AirPods Max, white bg
}


def slugify(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", str(text or "").lower())
    return re.sub(r"^-+|-+$", "", s)[:60]


def map_category(raw: str) -> str:
    key = str(raw or "Other").strip().lower()
    return CAT_MAP.get(key, "other")


ELEC_TITLE = re.compile(
    r"\b("
    r"iphone|ipad|airpods|macbook|galaxy\s*buds|earbuds?|earphone|headphone|headset|"
    r"playstation|ps5|xbox|game\s*console|power\s*bank|"
    r"jbl\s*electronics|samsung\s*electronics|apple\s*electronics|dyson\s*electronics|"
    r"electronics"
    r")\b",
    re.I,
)
ELEC_EXCLUDE = re.compile(
    r"\b(case|cover|bag|crossbody|camera\s*bag|t-?shirt|tee|hoodie|jacket|shoe|sneaker)\b",
    re.I,
)
HAT_TITLE = re.compile(r"\b(hat|beanie|cap)\b", re.I)

# Excel data-entry fixes (brand category wrong in source sheet)
CATEGORY_FORCE = {
    "7729045848": "accessories",  # Bape Hat tagged Electronics in Excel
}


def refine_category(item_id: str, title: str, category: str) -> str:
    """Fix obvious mislabels after Excel brand-category mapping."""
    if item_id in CATEGORY_FORCE:
        return CATEGORY_FORCE[item_id]
    t = title or ""
    if category == "electronics" and HAT_TITLE.search(t) and not ELEC_TITLE.search(t):
        return "accessories"
    if category != "electronics" and ELEC_TITLE.search(t) and not ELEC_EXCLUDE.search(t):
        return "electronics"
    return category


def seed_gallery(item_id: str) -> list[str]:
    paths: list[str] = []
    main = PRODUCTS_DIR / f"{item_id}.webp"
    if main.exists():
        paths.append(f"img/products/{item_id}.webp")
    for n in (2, 3, 4):
        alt = PRODUCTS_DIR / f"{item_id}-{n}.webp"
        if alt.exists():
            paths.append(f"img/products/{item_id}-{n}.webp")
    return paths


def money_usd(cny) -> float | None:
    try:
        v = float(cny)
    except (TypeError, ValueError):
        return None
    if v <= 0:
        return None
    return round(v / CNY_TO_USD, 2)


def build() -> dict:
    if not XLSX.exists():
        raise SystemExit(f"Missing sheet: {XLSX}")

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(h or "").strip() for h in next(ws.iter_rows(values_only=True))]
    idx = {h: i for i, h in enumerate(headers)}

    def cell(row, name, default=None):
        i = idx.get(name)
        if i is None or i >= len(row):
            return default
        return row[i]

    products = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        item_id = str(cell(row, "商品id") or "").strip()
        if not item_id or item_id in seen:
            continue
        seen.add(item_id)

        title = str(cell(row, "展示商品标题") or cell(row, "商品标题") or "").strip()
        if not title:
            continue

        gallery = seed_gallery(item_id)
        if not gallery:
            continue

        source_url = str(cell(row, "商品url") or "").strip()
        if not source_url:
            source_url = f"https://weidian.com/item.html?itemID={item_id}"

        platform = str(cell(row, "平台") or "weidian").strip().lower() or "weidian"
        brand = str(cell(row, "品牌") or "").strip()
        category = refine_category(item_id, title, map_category(cell(row, "品牌分类")))
        price = money_usd(cell(row, "商品优惠价"))
        score = cell(row, "商品最终得分")
        try:
            score_f = float(score) if score is not None else 0.0
        except (TypeError, ValueError):
            score_f = 0.0
        opens = cell(row, "打开商品详情页数")
        try:
            opens_n = int(opens or 0)
        except (TypeError, ValueError):
            opens_n = 0

        products.append(
            {
                "id": item_id,
                "platform": platform,
                "title": title,
                "brand": brand,
                "category": category,
                "price": price,
                "sourceUrl": source_url,
                "image": gallery[0],
                "gallery": gallery,
                # Future warehouse QC feed: fill batches here (photos[], weight_g, size_cm, created_at)
                "qc_batches": [],
                "score": round(score_f, 4),
                "opens": opens_n,
                "slug": f"{slugify(title)}-{item_id}" if slugify(title) else item_id,
            }
        )

    wb.close()

    products.sort(key=lambda p: (-p["opens"], -p["score"], p["title"]))

    # Top picks: high opens with multi-angle seed shots
    top_picks = [
        p["id"]
        for p in products
        if len(p["gallery"]) >= 2
    ][:24]

    trending = []
    for p in products:
        token = (p["brand"] or p["title"].split(" ")[0]).strip()
        if token and token not in trending:
            trending.append(token)
        if len(trending) >= 10:
            break

    cat_images = {}
    for p in products:
        cat_images.setdefault(p["category"], p["image"])

    categories = []
    for c in CATEGORIES:
        count = sum(1 for p in products if p["category"] == c["slug"])
        if not count:
            continue
        image = CAT_IMAGE_OVERRIDES.get(c["slug"]) or cat_images.get(c["slug"], products[0]["image"])
        categories.append(
            {
                **c,
                "count": count,
                "image": image,
            }
        )

    return {
        "generatedFrom": XLSX.name,
        "productCount": len(products),
        "qcFeedReady": True,
        "qcFeedNote": "qc_batches is empty until Kakobuy warehouse QC is authorized.",
        "categories": categories,
        "trending": trending,
        "topPickIds": top_picks,
        "products": products,
    }


def main() -> None:
    payload = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    body = (
        "window.GQ=window.GQ||{};\n"
        "GQ.catalog="
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        + ";\n"
    )
    OUT.write_text(body, encoding="utf-8")
    print(
        f"Wrote {payload['productCount']} products → {OUT.relative_to(ROOT)} "
        f"(top picks {len(payload['topPickIds'])})"
    )


if __name__ == "__main__":
    main()
